"""Regenerate FPR v1.7 DOCX from maintained source output/FPR_v1.4.md – full, not truncated."""
import re
from pathlib import Path
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

SOURCE = Path("output/FPR_v1.4.md")
TARGET = Path("24163800_Opaleye_Toluwalope_FPR_v1.7.docx")
FIG_ROOT = Path("output/figures")

def add_page_number(par):
    par.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = par.add_run()
    for tag, txt in [("begin",""),("instr"," PAGE "),("end","")]:
        el = OxmlElement(f"w:fldChar"); el.set(qn("w:fldCharType"), tag)
        if tag=="instr":
            el = OxmlElement("w:instrText"); el.set(qn("xml:space"),"preserve"); el.text=" PAGE "
        r._r.append(el)

def resolve_fig(name):
    for p in [FIG_ROOT/name, Path("output")/name, Path(name)]:
        if p.exists(): return p
    return None

doc = Document()
sec = doc.sections[0]
sec.top_margin = Inches(0.8); sec.bottom_margin = Inches(0.8)
sec.left_margin = Inches(1.0); sec.right_margin = Inches(1.0)
sec.header.paragraphs[0].text = "24163800 | Final Project Report | University of Hertfordshire"
add_page_number(sec.footer.paragraphs[0])

styles = doc.styles
styles["Normal"].font.name="Arial"; styles["Normal"].font.size=Pt(10.5)

lines = SOURCE.read_text(encoding="utf-8").splitlines()
# SSDF -> generic
lines = [l.replace("SSDF","secure development controls proposed").replace("SSDF controls proposed","secure development controls proposed") for l in lines]

i=0
while i < len(lines):
    t = lines[i]; i+=1
    if not t.strip(): continue
    # image markdown ![alt](path)
    m = re.search(r'!\[([^\]]*)\]\(([^)]+)\)', t)
    if m:
        cap, src = m.groups()
        src = src.strip()
        # resolve
        name = Path(src).name
        fig = resolve_fig(name)
        if fig and fig.exists():
            p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.add_run().add_picture(str(fig), width=Inches(5.9))
            cp = doc.add_paragraph(); cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r = cp.add_run(cap or name); r.italic=True; r.font.size=Pt(8.5)
        else:
            doc.add_paragraph(f"[Figure missing: {src}]")
        # also handle surrounding text
        before = t[:m.start()].strip()
        after = t[m.end():].strip()
        if before: doc.add_paragraph(before)
        if after: doc.add_paragraph(after)
        continue
    if t.startswith("|"):
        tbl_lines=[t]
        while i < len(lines) and lines[i].startswith("|"):
            tbl_lines.append(lines[i]); i+=1
        data=[]
        for tl in tbl_lines:
            cells=[c.strip() for c in tl.strip("|").split("|")]
            if all(set(c) <= {"-",":"} for c in cells): continue
            data.append(cells)
        if data:
            tbl = doc.add_table(rows=len(data), cols=len(data[0])); tbl.style="Table Grid"
            for r, row in enumerate(data):
                for c, val in enumerate(row):
                    tbl.cell(r,c).text=val
        continue
    if t.startswith("# "):
        p=doc.add_heading(t[2:], level=0); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    elif t.startswith("## "):
        doc.add_heading(t[3:], level=1)
    elif t.startswith("### "):
        doc.add_heading(t[4:], level=2)
    elif re.match(r"^\d+\. ", t):
        doc.add_paragraph(re.sub(r"^\d+\. ", "", t), style="List Number")
    elif t.startswith("- "):
        doc.add_paragraph(t[2:], style="List Bullet")
    elif t.startswith("> "):
        p=doc.add_paragraph(t[2:]); p.style="Intense Quote"
    elif t.startswith("**") and t.endswith("**"):
        p=doc.add_paragraph(); p.add_run(t.strip("*")).bold=True
    else:
        doc.add_paragraph(t.replace("`",""))

doc.save(TARGET)
print(f"Wrote {TARGET} ({TARGET.stat().st_size} bytes)")
# verify
import zipfile
z=zipfile.ZipFile(TARGET)
d=z.read("word/document.xml")
print(f"document.xml {len(d)} bytes, has Glossary? {'Glossary' in d.decode(errors='ignore')}, has SSDF? {'SSDF' in d.decode(errors='ignore')}")
