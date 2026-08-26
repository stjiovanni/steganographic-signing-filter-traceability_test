"""Embed captured dashboard screenshots as Figures 2-4 in FPR v1.5 -> v1.6,
replacing the '(image to be supplied by the author)' placeholder slots and
cleaning the List of Figures."""

import os
import shutil

from docx import Document
from docx.shared import Inches

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
SRC = os.path.join(ROOT, '24163800_Opaleye_Toluwalope_FPR_v1.5.docx')
TGT = os.path.join(ROOT, '24163800_Opaleye_Toluwalope_FPR_v1.6.docx')
FIGDIR = os.path.join(ROOT, 'output', 'figures')

if os.path.exists(TGT):
    os.remove(TGT)
shutil.copyfile(SRC, TGT)
d = Document(TGT)

figures = {
    'Figure 2': {
        'img': os.path.join(FIGDIR, 'dashboard_fig2_interface.png'),
        'caption': ('Figure 2: Dashboard Sign/Verify interface with the validated final1200 evidence set '
                    'active (dataset banner); Original and Signed preview panes shown after embedding.'),
    },
    'Figure 3': {
        'img': os.path.join(FIGDIR, 'dashboard_fig3_verify_single.png'),
        'caption': ('Figure 3: Single-method verification results (TrustMark): Detected = Yes, bit accuracy '
                    '100%, decoded payload TM00001.'),
    },
    'Figure 4': {
        'img': os.path.join(FIGDIR, 'dashboard_fig4_verify_hybrid.png'),
        'caption': ('Figure 4: Hybrid two-layer verification results: Two-layer recovery = Yes; TrustMark '
                    'TM00001 detected (100%); fallback channel FB01 recovered (98%).'),
    },
}

body_done = []
lof_done = []

for p in list(d.paragraphs):
    t = p.text.strip()
    if not t.startswith('Figure '):
        continue
    for key, spec in figures.items():
        if not t.startswith(key + ' (image to be supplied'):
            continue
        if 'CSV_DIR=' in t:
            # Body slot: insert image paragraph above, then rewrite caption
            pic_p = d.add_paragraph()
            pic_p.add_run().add_picture(spec['img'], width=Inches(6.0))
            p._p.addprevious(pic_p._p)
            new_caption = spec['caption']
            for r in list(p.runs):
                r._r.getparent().remove(r._r)
            p.add_run(new_caption)
            body_done.append(key)
        else:
            # List of Figures entry: strip the marker phrase
            new_text = t.replace('(image to be supplied by the author): ', ': ')
            for r in list(p.runs):
                r._r.getparent().remove(r._r)
            p.add_run(new_text)
            lof_done.append(key)

d.save(TGT)
print('Wrote', TGT)
print('body figures embedded:', sorted(set(body_done)))
print('LoF entries cleaned:', len(lof_done))