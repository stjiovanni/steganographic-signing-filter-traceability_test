"""Build FPR_Method_Comparison.xlsx: statistical comparison of the implemented
verification channels at final1200 scale.

Compared channels:
  * Hybrid two-layer (ours): TrustMark OR fallback-watermark decode
  * TrustMark alone: primary neural watermark
  * DCT baseline: strongest classical comparator
  * LSB baseline: fragile spatial reference (context)

Metrics (clearly separated, never conflated):
  * Payload recovery = exact decode-present rate over the 96,000 transformed
    observations (1,200 images x 80 conditions)
  * Clean-image recovery = same on the 1,200 untransformed images
  * Mean raw bit accuracy = pre-error-correction bit agreement (recovery is
    measured separately and is NOT derivable from this)

All numbers are computed from output/results/final1200/*.csv and
output/remediation_two_layer.json. Nothing estimated.
"""

import csv
import json
import os
from collections import defaultdict

from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.styles import Font, PatternFill

ROOT = os.path.dirname(os.path.abspath(__file__))
FINAL = os.path.join(ROOT, 'output', 'results', 'final1200')
OUT_PATH = os.path.join(ROOT, 'FPR_Method_Comparison.xlsx')

TRANSFORM_ORDER = ['brightness', 'contrast', 'saturation', 'vibrancy', 'gaussian_blur',
                   'salt_pepper_noise', 'jpeg_compression', 'rotation', 'scaling',
                   'crop_center', 'crop_random', 'letterbox']


def read_csv(directory, name):
    path = os.path.join(directory, name)
    with open(path, newline='', encoding='utf-8') as f:
        return list(csv.DictReader(f))


def recovery_by_transform(rows):
    """Exact decode-present rate per transform (transformed rows only)."""
    agg = defaultdict(lambda: [0, 0])
    for r in rows:
        if r['transform_name'] == 'none':
            continue
        agg[r['transform_name']][0] += 1
        agg[r['transform_name']][1] += r['decode_present'] == 'True'
    return {tf: round(ok / n, 4) for tf, (n, ok) in agg.items()}, \
           {tf: n for tf, (n, ok) in agg.items()}


def recovery_overall(rows):
    n = ok = 0
    clean_n = clean_ok = 0
    for r in rows:
        if r['transform_name'] == 'none':
            clean_n += 1
            clean_ok += r['decode_present'] == 'True'
        else:
            n += 1
            ok += r['decode_present'] == 'True'
    return round(ok / n, 4), round(clean_ok / clean_n, 4)


def bit_accuracy_overall(rows):
    vals = [float(r['bit_accuracy']) for r in rows if r['transform_name'] != 'none']
    return round(sum(vals) / len(vals), 4)


def main():
    tm = read_csv(FINAL, 'trustmark_robustness_results.csv')
    lsb = read_csv(FINAL, 'lsb_robustness_results.csv')
    dctm = read_csv(FINAL, 'dct_robustness_results.csv')

    with open(os.path.join(ROOT, 'output', 'remediation_two_layer.json'), encoding='utf-8') as f:
        TL = json.load(f)
    hybrid_tf = {t: TL['per_transform'][t]['two_layer'] for t in TRANSFORM_ORDER}
    hybrid_overall = TL['overall_transformed']['two_layer']
    hybrid_clean = TL['clean_baseline']['two_layer']

    tm_tf, ns = recovery_by_transform(tm)
    lsb_tf, _ = recovery_by_transform(lsb)
    dct_tf, _ = recovery_by_transform(dctm)
    tm_all, tm_clean = recovery_overall(tm)
    lsb_all, lsb_clean = recovery_overall(lsb)
    dct_all, dct_clean = recovery_overall(dctm)

    wb = Workbook()

    # ---- Sheet 1: Summary ----
    ws = wb.active
    ws.title = 'Summary'
    header = ['Channel', 'Payload recovery, transformed (%)', 'Clean-image recovery (%)',
              'Mean raw bit accuracy, transformed (%)', 'Observations per channel']
    ws.append(header)
    rows = [
        ['Hybrid two-layer (TrustMark OR fallback) - ours',
         round(hybrid_overall * 100, 2), round(hybrid_clean * 100, 2), None,
         TL['overall_transformed']['n']],
        ['TrustMark alone',
         round(tm_all * 100, 2), round(tm_clean * 100, 2),
         round(bit_accuracy_overall(tm) * 100, 2), len([r for r in tm if r['transform_name'] != 'none'])],
        ['DCT baseline',
         round(dct_all * 100, 2), round(dct_clean * 100, 2),
         round(bit_accuracy_overall(dctm) * 100, 2), len([r for r in dctm if r['transform_name'] != 'none'])],
        ['LSB baseline (context)',
         round(lsb_all * 100, 2), round(lsb_clean * 100, 2),
         round(bit_accuracy_overall(lsb) * 100, 2), len([r for r in lsb if r['transform_name'] != 'none'])],
    ]
    for row in rows:
        ws.append(row)
    note_row = len(rows) + 3
    ws.cell(row=note_row, column=1,
            value='Payload recovery = exact decode-present rate (ECC-corrected payload matches embedded '
                  'watermark). Mean raw bit accuracy is pre-error-correction bit agreement and is reported '
                  'separately; it does not imply recovery. Hybrid confidence/recovery definition: TrustMark '
                  'decode-present OR fallback decode-present. Scale: 1,200 MS-COCO images x 80 conditions.')
    ws.cell(row=note_row, column=1).alignment = __import__('openpyxl').styles.Alignment(wrap_text=True, vertical='top')
    ws.merge_cells(start_row=note_row, start_column=1, end_row=note_row + 2, end_column=5)

    # ---- Sheet 2: Per-transform recovery ----
    ws2 = wb.create_sheet('RecoveryByTransform')
    ws2.append(['Transform', 'Hybrid two-layer (ours)', 'TrustMark alone', 'DCT baseline', 'LSB baseline'])
    for tf in TRANSFORM_ORDER:
        ws2.append([tf,
                    round(hybrid_tf[tf] * 100, 2),
                    round(tm_tf[tf] * 100, 2),
                    round(dct_tf[tf] * 100, 2),
                    round(lsb_tf[tf] * 100, 2)])

    chart = BarChart()
    chart.type = 'col'
    chart.title = 'Payload recovery by transform (%, final1200)'
    chart.y_axis.title = 'Exact decode-present rate (%)'
    chart.x_axis.title = 'Transform'
    chart.height = 10
    chart.width = 24
    data = Reference(ws2, min_col=2, max_col=5, min_row=1, max_row=len(TRANSFORM_ORDER) + 1)
    cats = Reference(ws2, min_col=1, min_row=2, max_row=len(TRANSFORM_ORDER) + 1)
    chart.add_data(data, titles_from_data=True)
    chart.set_categories(cats)
    ws2.add_chart(chart, 'G2')

    # ---- Sheet 3: JPEG quality ladder ----
    ws3 = wb.create_sheet('JPEGQualityLadder')
    dct_jpeg = recovery_by_transform([r for r in dctm if r['transform_name'] == 'jpeg_compression'])[0]
    # per-quality needs intensities; recompute directly
    q_agg_dct = defaultdict(lambda: [0, 0])
    for r in dctm:
        if r['transform_name'] == 'jpeg_compression':
            q_agg_dct[r['intensity_value']][0] += 1
            q_agg_dct[r['intensity_value']][1] += r['decode_present'] == 'True'
    ws3.append(['JPEG quality', 'Hybrid two-layer (ours)', 'TrustMark alone', 'DCT baseline'])
    for row in TL['jpeg_quality_table']:
        q = row['quality']
        d = q_agg_dct.get(q, [1, 0])
        ws3.append([int(q),
                    round(row['two_layer'] * 100, 2),
                    round(row['trustmark'] * 100, 2),
                    round(d[1] / d[0] * 100, 2)])
    ws3.append(['Overall', round(TL['hypothesis']['jpeg_two_layer_recovery'] * 100, 2),
                round(TL['hypothesis']['jpeg_trustmark_only_recovery'] * 100, 2),
                round(sum(v[1] for v in q_agg_dct.values()) / sum(v[0] for v in q_agg_dct.values()) * 100, 2)])

    chart3 = BarChart()
    chart3.type = 'col'
    chart3.title = 'Payload recovery by JPEG quality (%, final1200)'
    chart3.y_axis.title = 'Exact decode-present rate (%)'
    chart3.x_axis.title = 'JPEG quality setting'
    chart3.height = 10
    chart3.width = 20
    data3 = Reference(ws3, min_col=2, max_col=4, min_row=1, max_row=9)
    cats3 = Reference(ws3, min_col=1, min_row=2, max_row=9)
    chart3.add_data(data3, titles_from_data=True)
    chart3.set_categories(cats3)
    ws3.add_chart(chart3, 'G2')

    # ---- Styling ----
    head_fill = PatternFill('solid', fgColor='1F4E78')
    for sheet in wb.worksheets:
        sheet.freeze_panes = 'A2'
        for cell in sheet[1]:
            cell.font = Font(bold=True, color='FFFFFF')
            cell.fill = head_fill
        for column in sheet.columns:
            width = max((len(str(c.value or '')) for c in column[:20]), default=8)
            sheet.column_dimensions[column[0].column_letter].width = min(max(width + 2, 10), 55)

    wb.save(OUT_PATH)
    print('Wrote', OUT_PATH)
    print('\n=== Summary ===')
    for row in rows:
        print(' ', row)


if __name__ == '__main__':
    main()