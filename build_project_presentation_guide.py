from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parent
FIGURES = ROOT / "output" / "figures"
OUT = ROOT / "PROJECT_PRESENTATION_GUIDE.docx"


def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_text(cell, text, bold=False, colour="172033"):
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(2)
    run = p.add_run(str(text))
    run.bold = bold
    run.font.name = "Aptos"
    run.font.size = Pt(8.5)
    run.font.color.rgb = RGBColor.from_string(colour)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def table(doc, headers, rows, widths=None):
    t = doc.add_table(rows=1, cols=len(headers))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.style = "Table Grid"
    for i, header in enumerate(headers):
        set_cell_text(t.rows[0].cells[i], header, bold=True, colour="FFFFFF")
        shade(t.rows[0].cells[i], "17324D")
    for row_index, row in enumerate(rows):
        cells = t.add_row().cells
        for i, value in enumerate(row):
            set_cell_text(cells[i], value)
            shade(cells[i], "F3F6F8" if row_index % 2 == 0 else "FFFFFF")
    if widths:
        for row in t.rows:
            for i, width in enumerate(widths):
                row.cells[i].width = Inches(width)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)
    return t


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("Page ")
    run.font.name = "Aptos"
    run.font.size = Pt(8)
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), "PAGE")
    paragraph._p.append(fld)


def add_figure(doc, filename, caption, width=6.1):
    path = FIGURES / filename
    if not path.exists():
        return
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(5)
    p.paragraph_format.space_after = Pt(3)
    p.add_run().add_picture(str(path), width=Inches(width))
    cp = doc.add_paragraph()
    cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cp.paragraph_format.space_after = Pt(8)
    r = cp.add_run(caption)
    r.italic = True
    r.font.name = "Aptos"
    r.font.size = Pt(8.5)
    r.font.color.rgb = RGBColor.from_string("425466")


def bullet(doc, text):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(3)
    p.add_run(text)
    return p


def note(doc, label, text):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.15)
    p.paragraph_format.right_indent = Inches(0.15)
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(7)
    r = p.add_run(label + " ")
    r.bold = True
    p.add_run(text)


def heading(doc, text, level=1):
    p = doc.add_heading(text, level=level)
    p.paragraph_format.keep_with_next = True
    return p


def page_break(doc):
    doc.add_page_break()


doc = Document()
sec = doc.sections[0]
sec.top_margin = Inches(0.65)
sec.bottom_margin = Inches(0.65)
sec.left_margin = Inches(0.78)
sec.right_margin = Inches(0.78)

styles = doc.styles
styles["Normal"].font.name = "Aptos"
styles["Normal"].font.size = Pt(10)
styles["Normal"].font.color.rgb = RGBColor.from_string("172033")
styles["Normal"].paragraph_format.space_after = Pt(6)
for name, size, colour in [("Title", 28, "17324D"), ("Heading 1", 19, "17324D"), ("Heading 2", 13, "0C6E72"), ("Heading 3", 10.5, "17324D")]:
    styles[name].font.name = "Aptos Display" if name == "Title" else "Aptos"
    styles[name].font.size = Pt(size)
    styles[name].font.bold = True
    styles[name].font.color.rgb = RGBColor.from_string(colour)

if "Caption" not in styles:
    styles.add_style("Caption", WD_STYLE_TYPE.PARAGRAPH)
styles["Caption"].font.name = "Aptos"
styles["Caption"].font.size = Pt(8.5)

header = sec.header.paragraphs[0]
header.text = "STEGANOGRAPHIC SIGNING AND FILTER TRACEABILITY"
header.runs[0].font.name = "Aptos"
header.runs[0].font.size = Pt(8)
header.runs[0].font.bold = True
header.runs[0].font.color.rgb = RGBColor.from_string("0C6E72")
add_page_number(sec.footer.paragraphs[0])

# Cover
p = doc.add_paragraph()
p.paragraph_format.space_before = Pt(52)
p.paragraph_format.space_after = Pt(14)
r = p.add_run("STUDY / PRESENTATION GUIDE")
r.bold = True
r.font.name = "Aptos"
r.font.size = Pt(11)
r.font.color.rgb = RGBColor.from_string("0C6E72")

sub = doc.add_paragraph()
sub.paragraph_format.space_before = Pt(5)
sub.paragraph_format.space_after = Pt(24)
sr = sub.add_run("A clear route through the robustness benchmark, its evidence, and its limits")
sr.font.size = Pt(14)
sr.font.color.rgb = RGBColor.from_string("425466")

table(doc, ["Project", "Evidence baseline"], [
    ("MSc Computer Science project", "Validated final1200 evidence set"),
    ("Student", "Opaleye Toluwalope Abayomi"),
    ("Student number", "24163800"),
    ("Core question", "Which verification signals remain useful after controlled image transformations?"),
], widths=[1.7, 4.5])
note(doc, "Reading rule:", "Treat every result as conditional on the implementation, dataset, transform, metric, and validation state described in this guide. A recovered signal is not proof of truth, identity, or provenance.")
add_figure(doc, "fpr_transform_taxonomy.png", "Figure 1. Transform taxonomy used to organise the robustness study. Source: output/figures/fpr_transform_taxonomy.png.", 5.9)

page_break(doc)
heading(doc, "1. Executive Overview", 1)
doc.add_paragraph("This project benchmarks complementary image-verification signals rather than searching for one universal authenticity detector. It compares perceptual hashes, a neural watermark, and two classical watermark baselines under a shared matrix of graduated image transformations. The study is designed to show where each signal remains informative and where it fails.")
heading(doc, "The central message", 2)
doc.add_paragraph("Robustness is conditional. TrustMark is generally strong against value and encoding changes in the tested evidence, while geometric changes such as rotation, cropping, and letterbox padding expose important weaknesses. Perceptual hashes provide a separate similarity signal; LSB and DCT provide useful baselines but are not production recommendations. The ensemble is a decision aid, not a trust system.")
table(doc, ["Question", "Answer supported by the study"], [
    ("What is being measured?", "Persistence of a verification signal after a defined transformation."),
    ("What is not being measured?", "Authorship, factual truth, identity, signed provenance, or resistance to adaptive attackers."),
    ("Why use several methods?", "Different embedding domains and comparison rules fail under different transformations."),
    ("What is the practical output?", "Auditable result rows, threshold analysis, an ensemble matrix, and an inspectable dashboard."),
], widths=[1.65, 4.55])
heading(doc, "Suggested opening", 2)
note(doc, "Thirty-second version:", "Images are routinely resized, recompressed, filtered, cropped, and re-encoded. This project asks whether several verification signals degrade gracefully under those operations. It uses a shared deterministic pipeline, records row-level evidence, and reports the boundary between useful similarity evidence and unsupported claims of authenticity.")
heading(doc, "Three claims to remember", 2)
bullet(doc, "A low perceptual-hash distance indicates similarity under a chosen rule; it is not cryptographic authentication.")
bullet(doc, "A decoded watermark demonstrates signal recovery for the tested payload and condition; it is not proof of origin.")
bullet(doc, "The strongest contribution is traceability: explicit transforms, validation checks, corrections, limitations, and reproducible evidence paths.")

page_break(doc)
heading(doc, "2. Study Design", 1)
doc.add_paragraph("The experiment applies each transform independently to the same deterministically selected image set. This isolates the effect of a named operation and makes method-to-method comparisons easier to interpret. It does not model cascaded attacks or adaptive adversaries.")
table(doc, ["Element", "Design"], [
    ("Dataset", "MS-COCO 2017 validation images; deterministic filename-order selection."),
    ("Evidence scale", "Validated final1200 run; image is the statistical unit for uncertainty summaries."),
    ("Transform matrix", "12 transform families and 80 transform-intensity conditions."),
    ("Hash layer", "Eight perceptual-hash variants; reference-to-transformed Hamming distance."),
    ("Watermark layer", "TrustMark, LSB, and differential DCT baselines with fixed payloads."),
    ("Analysis", "Threshold crossings, worst-algorithm hash rule, and best/fallback ensemble choices."),
], widths=[1.55, 4.65])
heading(doc, "Transform families", 2)
table(doc, ["Family", "Examples", "Interpretive focus"], [
    ("Photometric", "Brightness, contrast, saturation, vibrancy", "Pixel-value changes"),
    ("Encoding / noise", "Blur, salt-and-pepper noise, JPEG", "Filtering and re-encoding"),
    ("Geometric", "Rotation, scaling, centre crop, random crop", "Coordinate and content disruption"),
    ("Layout", "Black and grey letterbox padding", "Aspect-ratio and padding effects"),
], widths=[1.35, 2.45, 2.4])
add_figure(doc, "fpr_method_comparison.png", "Figure 2. Method-comparison view used to frame the complementary roles of hashes and watermarking. Source: output/figures/fpr_method_comparison.png.", 5.85)

page_break(doc)
heading(doc, "3. Signals and Metrics", 1)
heading(doc, "Perceptual hashes", 2)
doc.add_paragraph("The hash pipeline computes a reference hash for the original image and a second hash after each transformation. Hamming distance counts differing bits. Because hash lengths differ, the analysis also uses a bit-error fraction. PDQ is represented as a 256-bit value; the remaining configured hashes use their implementation-specific lengths.")
heading(doc, "Watermarks", 2)
doc.add_paragraph("TrustMark embeds a fixed payload with a learned encoder and decoder. LSB writes a payload into the red-channel least-significant bit plane. DCT differentially embeds payload bits in mid-frequency coefficients. Their shared purpose is comparison, not equivalence: payload construction, embedding mechanism, decoder, and failure semantics differ.")
table(doc, ["Metric", "Meaning", "Do not confuse it with"], [
    ("Hamming distance", "Number of changed hash bits.", "A cryptographic proof of identity."),
    ("Bit-error fraction", "Hamming distance divided by the relevant hash length.", "A universal image-severity scale."),
    ("Raw bit accuracy", "Recovered payload bits before error correction.", "Exact payload recovery."),
    ("Decode presence", "Whether the implementation reports a valid decoded payload.", "A calibrated probability of authenticity."),
    ("PSNR / MSE", "Encode quality or distortion measure.", "Robustness or security."),
], widths=[1.35, 2.9, 2.0])
heading(doc, "Decision rules", 2)
doc.add_paragraph("Watermark success uses mean bit accuracy above 0.5. Hash success uses a mean bit-error fraction below 0.5. For the ensemble, the hash signal uses the worst-algorithm mean so that a group of stable 64-bit hashes cannot hide a weak member such as PDQ. The matrix records a best method and a fallback for each condition.")
note(doc, "Important:", "The 0.5 boundary is an explicit analysis convention. It is not a security proof, a user-validated operating point, or a substitute for false-acceptance and false-rejection analysis.")

page_break(doc)
heading(doc, "4. Findings to Present", 1)
heading(doc, "Observed pattern", 2)
doc.add_paragraph("The validated final-run threshold analysis shows a clear split between content-preserving changes and spatial reorganisation. Hash and TrustMark thresholds are not crossed for most tested families, whereas the classical baselines cross the 50% boundary in several important conditions.")
table(doc, ["Condition", "Evidence-led talking point"], [
    ("Brightness / contrast", "LSB crosses the 50% boundary at selected mild-to-moderate settings; hash and TrustMark remain below their critical thresholds in the reported table."),
    ("JPEG compression", "LSB crosses at quality 20; DCT remains the more informative classical comparator in this study."),
    ("Rotation", "DCT crosses at 1 degree; the worst-hash rule crosses at 70 degrees; LSB crosses at 90 degrees."),
    ("Scaling", "LSB crosses at 0.25x; DCT crosses at 0.25x, with a non-monotonic worst crossing recorded at 3.0x."),
    ("Centre crop", "Worst-hash crossing occurs at 0.7 removal; LSB and DCT cross earlier under the tested settings."),
    ("Random crop", "Worst-hash crossing occurs at 0.6 removal; LSB and DCT cross at lower tested retention levels."),
    ("Letterbox", "Padding changes layout without removing the source content, making it a useful condition-specific comparison rather than a universal ranking test."),
], widths=[1.55, 4.65])
add_figure(doc, "ensemble_heatmap.png", "Figure 3. Ensemble heatmap showing method choices across transform-intensity conditions. Source: output/figures/ensemble_heatmap.png.", 5.85)

page_break(doc)
heading(doc, "5. Evidence Walkthrough", 1)
heading(doc, "Recommended presentation sequence", 2)
table(doc, ["Stage", "Show / say", "Purpose"], [
    ("1. Problem", "Show the transform taxonomy; explain routine edits and geometric disruption.", "Establish why robustness matters."),
    ("2. Design", "Show the shared pipeline and four signal families.", "Make the comparison fair and auditable."),
    ("3. Metrics", "Explain Hamming distance, raw bit accuracy, decode presence, and PSNR.", "Prevent metric overclaiming."),
    ("4. Results", "Use the method comparison and heatmap to contrast failure regions.", "Make complementarity visible."),
    ("5. Dashboard", "Open overview, per-transform, ensemble, and per-image views.", "Move from aggregate pattern to row-level evidence."),
    ("6. Boundary", "State what the results do not establish.", "Close with defensible interpretation."),
], widths=[0.9, 3.25, 2.05])
heading(doc, "Dashboard demonstration", 2)
doc.add_paragraph("The dashboard is a research visualisation backed by CSV result files. The overview summarises coverage; the per-transform view shows method behaviour against intensity; the ensemble matrix exposes best and fallback choices; and the per-image view supports row-level inspection. It should be described as an evidence browser, not as a production signing service.")
add_figure(doc, "dashboard_fig2_interface.png", "Figure 4. Dashboard interface for navigating the benchmark evidence. Source: output/figures/dashboard_fig2_interface.png.", 5.75)
add_figure(doc, "dashboard_fig4_verify_hybrid.png", "Figure 5. Hybrid verification view illustrating complementary signal inspection. Source: output/figures/dashboard_fig4_verify_hybrid.png.", 5.75)

page_break(doc)
heading(doc, "6. Quality, Corrections, and Reproducibility", 1)
heading(doc, "Quality controls", 2)
bullet(doc, "The transform definitions are centralised in `transforms.py`, reducing configuration drift between pipelines.")
bullet(doc, "Image selection and random transform seeds are deterministic and recorded in project evidence.")
bullet(doc, "Validation checks coverage, expected rows, duplicate keys, missing values, and baseline records before interpretation.")
bullet(doc, "Raw CSV outputs are retained separately from summary tables and dashboard views.")
bullet(doc, "Uncertainty summaries use image-level aggregation, a fixed-seed percentile bootstrap for continuous metrics, and Wilson intervals for decode proportions.")
heading(doc, "Corrections that matter", 2)
table(doc, ["Correction", "Why it changes interpretation"], [
    ("PDQ serialisation", "Each dimension is represented as one bit rather than an incorrect multi-character representation."),
    ("TrustMark metric", "Raw decoder output before error correction is measured separately from exact decoded-payload presence."),
    ("Sample-size boundary", "Configured inventory and runners are not treated as completed evidence until all pipelines pass validation."),
], widths=[1.55, 4.65])
heading(doc, "Reproducibility checklist", 2)
table(doc, ["Check", "Evidence to retain"], [
    ("Inputs", "Image manifest, dataset description, and checksums where practicable."),
    ("Environment", "Package versions, model version, operating environment, and pipeline version."),
    ("Execution", "Run log, command arguments, output paths, and validation result."),
    ("Analysis", "Generating scripts, threshold definitions, aggregation rules, and figure sources."),
    ("Interpretation", "Explicit sample size, uncertainty method, and limitations in every result context."),
], widths=[1.35, 4.85])
add_figure(doc, "fpr_payload_ecc_dev_sweep.png", "Figure 6. Payload and error-correction development sweep used to motivate configuration choices. Source: output/figures/fpr_payload_ecc_dev_sweep.png.", 5.8)

page_break(doc)
heading(doc, "7. Limitations and Responsible Use", 1)
heading(doc, "Current limitations", 2)
bullet(doc, "The evidence is tied to the selected MS-COCO images and does not establish performance on medical, synthetic, text-heavy, or customer imagery.")
bullet(doc, "The transformations are isolated. Compound, adaptive, and adversarial attacks are outside the demonstrated scope.")
bullet(doc, "Only the configured TrustMark implementation and payload are evaluated; other models, payload sizes, and keys may behave differently.")
bullet(doc, "The ensemble matrix reports thresholded comparative evidence, not calibrated authentication accuracy or decision cost.")
bullet(doc, "The dashboard is CSV-backed. A migration script does not demonstrate a deployed, secured database service.")
bullet(doc, "C2PA manifests, certificates, signed claims, key management, identity, and a trust chain are not implemented by this benchmark.")
heading(doc, "Ethics and professional responsibility", 2)
doc.add_paragraph("The study uses a public research image dataset and does not recruit participants or collect new personal data. Public availability does not remove rights, privacy, or misuse considerations: images may depict recognisable people or copyrighted material. Future payloads containing identifiers, locations, or timestamps would require a documented purpose, lawful basis, retention policy, access control, and deletion process.")
heading(doc, "Safe wording", 2)
table(doc, ["Prefer", "Avoid"], [
    ("The method retained useful signal under the tested condition.", "The method proves the image is authentic."),
    ("The result is descriptive for this validated sample.", "The result generalises to all digital images."),
    ("The dashboard exposes comparative evidence.", "The dashboard is a production provenance service."),
    ("The watermark is a soft verification signal.", "The watermark is a cryptographic signature."),
], widths=[3.1, 3.1])

page_break(doc)
heading(doc, "8. Questions and Short Answers", 1)
qa = [
    ("Why not use only a watermark?", "A watermark and a hash measure different relationships. A hybrid view can preserve useful evidence when one signal degrades, although it still needs calibrated decisions and provenance controls."),
    ("Why is geometry difficult?", "Rotation, crop, scaling, and padding alter coordinate relationships. A signal trained or designed for value changes may not retain the same relationship after spatial reorganisation."),
    ("Does a high bit accuracy mean the watermark is secure?", "No. It means the tested decoder recovered the payload accurately under the tested condition. Security also requires a threat model, key management, attack evaluation, and misuse controls."),
    ("Why report raw bit accuracy and decode presence?", "Error correction can make exact decoding succeed even when some raw bits are wrong. Reporting both reveals that distinction."),
    ("Why use the worst hash in the ensemble?", "A pooled mean can hide a weak algorithm, particularly when most hashes have shorter representations. The weakest-link rule makes that choice explicit."),
    ("What would strengthen the study next?", "Complete validated larger-run evidence, add compound and adaptive attacks, report per-image uncertainty and error costs, and evaluate provenance and security controls separately."),
]
for question, answer in qa:
    heading(doc, question, 2)
    doc.add_paragraph(answer)

heading(doc, "One-minute closing", 2)
doc.add_paragraph("The benchmark shows that verification signals have different failure surfaces. TrustMark, perceptual hashes, LSB, and DCT are therefore best understood as conditional evidence layers, not interchangeable proofs. The project's engineering value is the shared deterministic pipeline, the inspectable evidence trail, and the discipline to keep robustness findings separate from claims of authenticity or provenance.")

page_break(doc)
heading(doc, "Appendix A. Source Map", 1)
doc.add_paragraph("This guide was assembled from the repository materials available at generation time. The source map makes the document auditable without modifying the FPR or dashboard code.")
table(doc, ["Topic", "Repository source"], [
    ("Project scope and pipeline components", "PROJECT_SUMMARY.md; README.md"),
    ("Method definitions and limitations", "output/final_project_report_draft.md"),
    ("Final-run uncertainty method", "output/results/final1200/uncertainty_summary.md"),
    ("Threshold crossings", "output/results/final1200/threshold_analysis.md"),
    ("Figures", "output/figures/*.png; captions identify each selected asset"),
    ("Evidence and submission context", "SUBMISSION_MANIFEST.md; CITATIONS.md"),
], widths=[2.2, 4.0])
heading(doc, "Figure index", 2)
table(doc, ["Figure", "Asset"], [
    ("1", "fpr_transform_taxonomy.png"),
    ("2", "fpr_method_comparison.png"),
    ("3", "ensemble_heatmap.png"),
    ("4", "dashboard_fig2_interface.png"),
    ("5", "dashboard_fig4_verify_hybrid.png"),
    ("6", "fpr_payload_ecc_dev_sweep.png"),
], widths=[0.8, 5.4])
note(doc, "Status:", "This is a study and presentation document derived from repository evidence. It is not a replacement for the formal project report, institutional declarations, or a signed provenance implementation.")

page_break(doc)
heading(doc, "Appendix B. Glossary", 1)
doc.add_paragraph("Terms as used in this guide and the final report (22 entries).")
table(doc, ["Term", "Definition"], [
    ("Authenticity", "Claim that content is true or correctly attributed; not established by signal recovery alone."),
    ("BCH", "Bose-Chaudhuri-Hocquenghem error-correcting code used in watermark packets."),
    ("Bit accuracy", "Proportion of expected watermark bits recovered before error correction."),
    ("Decode presence", "Whether the decoder reports a valid payload; distinct from raw bit accuracy."),
    ("DCT", "Discrete cosine transform; frequency-domain baseline embedding in mid-frequency coefficients."),
    ("ECC", "Error-correcting code; converts partial bit recovery into valid messages."),
    ("Ensemble", "Decision matrix combining hash and watermark signals per condition."),
    ("False acceptance", "Accepting an unrelated or incorrectly watermarked image as verified."),
    ("Hamming distance", "Number of differing bits between reference and transformed hash."),
    ("Integrity", "Evidence that bytes or content are unchanged; requires hard binding."),
    ("Letterbox", "Padding added to preserve aspect ratio; separate from cropping."),
    ("LSB", "Least-significant-bit embedding; fragile spatial-domain baseline."),
    ("MSE", "Mean squared error; pixel-domain distortion measure."),
    ("Payload", "Fixed test message embedded and recovered; not an identifier."),
    ("Perceptual hash", "Compact content-derived representation compared by distance."),
    ("pHash/dHash", "Specific perceptual hash variants (among eight tested)."),
    ("Provenance", "Signed, accountable record of claims; requires manifests and trust chain."),
    ("PSNR", "Peak signal-to-noise ratio; imperceptibility metric, not robustness."),
    ("Recovery", "Whether a payload decodes; see decode presence."),
    ("TrustMark", "Learned neural watermark method used as primary candidate."),
    ("Two-layer fallback", "Second BCH-coded DCT layer; recovery is TrustMark OR fallback."),
    ("Transform intensity", "Parameter value of a named condition; comparable only within family."),
], widths=[1.65, 4.55])

doc.save(OUT)
print(OUT)
