from docx import Document
from docx.oxml.ns import qn

d = Document('24163800_Opaleye_Toluwalope_FPR_v1.3.docx')
for p in d.paragraphs:
    if 'CSV-backed single-page web application' in p.text or 'dashboard is a' in p.text.lower():
        print('DASH:', p.text[:400])
        break
# Confirm the screenshot slots reference Sign/Verify
print('--- figure screenshot slots ---')
for p in d.paragraphs:
    if 'image to be supplied' in p.text:
        print(' ', p.text[:130])
