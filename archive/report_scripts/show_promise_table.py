from docx import Document
from docx.oxml.ns import qn
# Print the promise-vs-delivery table (Table 2) content from FPR v1.3
d = Document('24163800_Opaleye_Toluwalope_FPR_v1.3.docx')
# Table index 1 should be the promise table (T1 objectives=0, T2 promise=1)
if len(d.tables) > 1:
    t = d.tables[1]
    print('=== Promise-vs-delivery table (index 1) ===')
    for r in t.rows:
        print(' | '.join(c.text[:70] for c in r.cells))
