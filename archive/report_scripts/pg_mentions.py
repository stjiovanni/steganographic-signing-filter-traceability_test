from docx import Document

for path, label in [('24163800_OpaleyeToluwalope_DPP_v2.docx', 'DPP'),
                    ('24163800_Tolu_Opaleye_IPR_Report_V4.docx', 'IPR')]:
    d = Document(path)
    print('==== ' + label + ' ====')
    for i, p in enumerate(d.paragraphs):
        t = p.text.strip()
        low = t.lower()
        if any(k in low for k in ('postgres', 'database', 'sql', 'persistence')):
            print(i, '|', t[:300])
    for ti, t in enumerate(d.tables):
        for r in t.rows:
            joined = ' || '.join(c.text for c in r.cells)
            if any(k in joined.lower() for k in ('postgres', 'database', 'sql', 'persistence')):
                print('TABLE', ti, '|', joined[:200])