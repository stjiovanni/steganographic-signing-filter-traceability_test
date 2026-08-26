from docx import Document
d = Document('24163800_Opaleye_Toluwalope_FPR_v1.3.docx')
for p in d.paragraphs:
    if 'Table 14 reports recovery by JPEG quality' in p.text:
        new_text = p.text.replace('Table 14 reports recovery by JPEG quality.',
                                  'Table 12 reports recovery by JPEG quality.')
        for run in list(p.runs):
            run._r.getparent().remove(run._r)
        p.add_run(new_text)
        print('fixed cross-ref')
        break
d.save('24163800_Opaleye_Toluwalope_FPR_v1.3.docx')
print('saved')