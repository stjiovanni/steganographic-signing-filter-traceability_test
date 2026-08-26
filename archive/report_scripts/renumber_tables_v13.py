"""Renumber all table captions in FPR v1.3 in document order and update the
List of Tables. The new promise-vs-delivery table in Section 1 shifts every
later table number by one.

Document order of tables (by first-appearance of their caption):
  §1.3 objectives -> Table 1
  §1.5 promise-vs-delivery (new) -> Table 2
  §3.2 evidence set -> Table 3
  §3.5 metrics -> Table 4
  §3.6 gates -> Table 5
  §3.6 gate audit -> Table 6
  §4.2 evidence matrix -> Table 7
  §4.6 final1200 means -> Table 8
  §4.7 ECC ranking -> Table 9
  §4.10 uncertainty -> Table 10
  §4.11 two-layer -> Table 11
  §4.11 JPEG quality -> Table 12
  §5.4 risk -> Table 13
  §5.5 Gantt -> Table 14
  §5.6 BCS -> Table 15
"""

from docx import Document
from docx.oxml.ns import qn
import os
import shutil

SRC = '24163800_Opaleye_Toluwalope_FPR_v1.3.docx'
TGT = '24163800_Opaleye_Toluwalope_FPR_v1.3.docx'
d = Document(SRC)

# Map of old caption keyword -> new table number
# Identify each table caption paragraph by a unique substring and set its number.
caption_map = [
    ('Objectives and success criteria', 1),
    ('IPR/DPP promise versus delivery', 2),
    ('Evidence sets, artefacts and permitted use', 3),
    ('Metric definitions and their evidentiary limits', 4),
    ('Validation gates and failure responses', 5),
    ('Validation gate enforcement audit', 6),
    ('Evidence matrix: claims, sources, scale and permitted wording', 7),
    ('Final1200 condition-level means by method and transform', 8),
    ('Payload/ECC ablation ranking (12 configurations', 9),
    ('Final1200 image-level uncertainty summary', 10),
    ('Final1200 two-layer payload-recovery rate by transform', 11),
    ('Final1200 payload recovery by JPEG quality', 12),
    ('Threat and risk matrix', 13),
    ('Project work packages and outcomes', 14),
    ('BCS Code of Conduct mapping', 15),
]


def set_caption_number(paragraph, number):
    text = paragraph.text
    # Replace "Table N:" at the start with "Table {number}:"
    import re
    new_text = re.sub(r'^Table \d+:', f'Table {number}:', text)
    if new_text != text:
        for run in list(paragraph.runs):
            run._r.getparent().remove(run._r)
        paragraph.add_run(new_text)
        return True
    return False


changed = 0
for p in d.paragraphs:
    t = p.text.strip()
    if not t.startswith('Table '):
        continue
    for keyword, number in caption_map:
        if keyword in t:
            if set_caption_number(p, number):
                changed += 1
            break

# Update List of Tables entries
# The List of Tables section has entries like "Table N: <title> (Section ...)".
list_map = [
    ('Objectives and success criteria (Section 1.3)', 1),
    ('Promise versus delivery (Section 1.5)', 2),
    ('Evidence sets, artefacts and permitted use (Section 3.2)', 3),
    ('Metric definitions and their evidentiary limits (Section 3.5)', 4),
    ('Validation gates and failure responses (Section 3.6)', 5),
    ('Validation gate enforcement audit (Section 3.6)', 6),
    ('Evidence matrix: claims, sources, scale and permitted wording (Section 4.2)', 7),
    ('Final1200 condition-level means by method and transform (Section 4.6)', 8),
    ('Payload/ECC ablation development-scale ranking (Section 4.7)', 9),
    ('Final1200 image-level uncertainty summary (Section 4.10)', 10),
    ('Final1200 two-layer payload-recovery rate by transform (Section 4.11)', 11),
    ('Final1200 payload recovery by JPEG quality (Section 4.11)', 12),
    ('Threat and risk matrix (Section 5.4)', 13),
    ('Project work packages and outcomes (Section 5.5)', 14),
    ('BCS Code of Conduct mapping (Section 5.6)', 15),
]

# Rebuild List of Tables body. Find the paragraph list between "List of tables" heading
# and "Glossary" heading.
body = d.element.body
seq = []
for c in body.iterchildren():
    seq.append(c)

# Find List of tables heading and Glossary heading indexes
lt_idx = None
gl_idx = None
for i, c in enumerate(seq):
    tag = c.tag.split('}')[1]
    if tag != 'p':
        continue
    txt = ''.join(t.text or '' for t in c.iter(qn('w:t'))).strip()
    if txt == 'List of tables':
        lt_idx = i
    if txt == 'Glossary':
        gl_idx = i
        break

# Remove existing list-of-tables entries (paragraphs between lt_idx and gl_idx)
if lt_idx is not None and gl_idx is not None:
    to_remove = [seq[i] for i in range(lt_idx + 1, gl_idx)]
    for el in to_remove:
        el.getparent().remove(el)
    # Insert new entries after the "List of tables" heading
    anchor = seq[lt_idx]
    from docx.text.paragraph import Paragraph
    from docx.oxml import OxmlElement
    for label, number in list_map:
        new_p = OxmlElement('w:p')
        anchor.addnext(new_p)
        p = Paragraph(new_p, d.paragraphs[0]._parent)
        p.add_run(f'Table {number}: {label}.')
        anchor = new_p

d.save(TGT)
print('caption renumber changed:', changed)
print('saved', TGT)