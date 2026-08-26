"""Convert the maintained Markdown FPR draft into a submission-editable DOCX."""

import argparse
import re
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


SOURCE = 'output/FPR_v0.5.md'
TARGET = '24163800_Opaleye_Toluwalope_FPR_v0.5.docx'


def add_page_number(paragraph):
    run = paragraph.add_run()
    begin = OxmlElement('w:fldChar')
    begin.set(qn('w:fldCharType'), 'begin')
    instruction = OxmlElement('w:instrText')
    instruction.set(qn('xml:space'), 'preserve')
    instruction.text = ' PAGE '
    end = OxmlElement('w:fldChar')
    end.set(qn('w:fldCharType'), 'end')
    run._r.extend([begin, instruction, end])


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument('--source', default=SOURCE)
    parser.add_argument('--target', default=TARGET)
    args = parser.parse_args()
    source = args.source
    target = args.target

    document = Document()
    section = document.sections[0]
    section.top_margin = Inches(0.8)
    section.bottom_margin = Inches(0.8)
    section.left_margin = Inches(1.0)
    section.right_margin = Inches(1.0)
    styles = document.styles
    styles['Normal'].font.name = 'Arial'
    styles['Normal'].font.size = Pt(10.5)
    lines = open(source, encoding='utf-8').read().splitlines()
    index = 0
    while index < len(lines):
        text = lines[index]
        index += 1
        if not text.strip():
            continue
        if text.startswith('|'):
            table_lines = [text]
            while index < len(lines) and lines[index].startswith('|'):
                table_lines.append(lines[index])
                index += 1
            data = []
            for table_line in table_lines:
                cells = [cell.strip() for cell in table_line.strip('|').split('|')]
                if all(set(cell) <= {'-', ':'} for cell in cells):
                    continue
                data.append(cells)
            if data:
                table = document.add_table(rows=len(data), cols=len(data[0]))
                table.style = 'Table Grid'
                for row_index, row in enumerate(data):
                    for col_index, value in enumerate(row):
                        table.cell(row_index, col_index).text = value
            continue
        if text.startswith('# '):
            p = document.add_heading(text[2:], level=0)
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        elif text.startswith('## '):
            document.add_heading(text[3:], level=1)
        elif text.startswith('### '):
            document.add_heading(text[4:], level=2)
        elif re.match(r'^\d+\. ', text):
            document.add_paragraph(re.sub(r'^\d+\. ', '', text), style='List Number')
        elif text.startswith('- '):
            document.add_paragraph(text[2:], style='List Bullet')
        elif text.startswith('> '):
            p = document.add_paragraph(text[2:])
            p.style = 'Intense Quote'
        elif text.startswith('**') and text.endswith('**'):
            p = document.add_paragraph()
            p.add_run(text.strip('*')).bold = True
        else:
            document.add_paragraph(text.replace('`', ''))
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.add_run('24163800 | Final Project Report | ')
    add_page_number(footer)
    document.save(target)
    print(f'Wrote {target}')


if __name__ == '__main__':
    main()
