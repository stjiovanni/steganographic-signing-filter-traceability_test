from docx import Document

DOC_PATH = r"C:\Users\hp\Documents\porfolio\msc_proj\IPR_Report_V1.docx"

doc = Document(DOC_PATH)

def replace_para_text(para, new_text):
    """Clear all runs and set new text in the first run, preserving the first run's formatting."""
    for run in para.runs:
        run.text = ""
    para.runs[0].text = new_text

def find_para_index(doc, substr):
    for i, p in enumerate(doc.paragraphs):
        if substr in p.text:
            return i
    return None

# --- Fix 1: Tone down Mareen et al. claim (Section 1.2, para 15) ---
idx = find_para_index(doc, "Mareen et al. (2021) argue explicitly for perceptual hashing as a fast fallback")
if idx is not None:
    old = doc.paragraphs[idx].text
    new = old.replace(
        "Mareen et al. (2021) argue explicitly for perceptual hashing as a fast fallback mechanism to be used alongside a primary watermark, rather than as a replacement for one, since perceptual hashes can be computed and compared far faster than a full watermark verification pass.",
        "Mareen et al. (2021) propose a fallback watermark detection system using secondary watermarks for faster comparison, an architectural pattern that parallels this project's use of perceptual hashing as a lightweight complement to the primary TrustMark watermark."
    )
    replace_para_text(doc.paragraphs[idx], new)
    print(f"[Fix 1] Para {idx}: Mareen et al. claim toned down.")
else:
    print("[Fix 1] ERROR: Could not find target text.")

# --- Fix 2: Remove Singhi et al. placeholder (Section 1.2, para 18) ---
idx = find_para_index(doc, "[INSERT: your DPP-cited Singhi et al. source on provenance detection")
if idx is not None:
    new_text = (
        "A relevant recent preprint by Singhi et al. (2025) proposes DinoHash, an adversarially robust "
        "perceptual hash derived from DINOv2, combined with homomorphic encryption for provenance detection "
        "of AI-generated images. While this preprint has not yet undergone peer review, its architectural "
        "combination of perceptual hashing with cryptographic verification aligns with this project's hybrid "
        "design philosophy and represents a promising direction for future work (see Section 3.7)."
    )
    replace_para_text(doc.paragraphs[idx], new_text)
    print(f"[Fix 2] Para {idx}: Singhi et al. placeholder replaced.")
else:
    print("[Fix 2] ERROR: Could not find target text.")

# --- Fix 3: Update reference list - Singhi citation (Section 6, para 83) ---
idx = find_para_index(doc, "[INSERT: your Singhi et al. citation")
if idx is not None:
    new_text = (
        "Singhi, S., Yadav, A., Gupta, A., Ebrahimi, S.M. and Hassanizadeh, P. (2025) "
        "'Provenance Detection for AI-Generated Images: Combining Perceptual Hashing, Homomorphic "
        "Encryption, and AI Detection Models', arXiv:2503.11195. Available at: "
        "https://arxiv.org/abs/2503.11195 (Accessed: 12 July 2026)."
    )
    replace_para_text(doc.paragraphs[idx], new_text)
    print(f"[Fix 3] Para {idx}: Singhi citation replaced.")
else:
    print("[Fix 3] ERROR: Could not find target text.")

# Verify TrustMark URL is correct (para 78)
idx = find_para_index(doc, "github.com/adobe/trustmark")
if idx is not None:
    print(f"[Fix 3] Para {idx}: TrustMark URL already correct (adobe/trustmark).")
else:
    # Check if wrong URL exists
    idx2 = find_para_index(doc, "github.com/benibaez/trustmark")
    if idx2 is not None:
        old = doc.paragraphs[idx2].text
        new = old.replace("benibaez/trustmark", "adobe/trustmark")
        replace_para_text(doc.paragraphs[idx2], new)
        print(f"[Fix 3] Para {idx2}: TrustMark URL corrected from benibaez to adobe.")
    else:
        print("[Fix 3] WARNING: Could not find TrustMark reference URL to verify.")

# --- Fix 4: Add dashboard screenshot placeholder in Appendix 2 (after para 94) ---
idx = find_para_index(doc, "[INSERT: original vs. transformed image comparisons")
if idx is not None:
    # We need to add a new paragraph after this one. We'll insert by adding a paragraph
    # to the end of the document body and then moving it. Simpler: use XML manipulation.
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement
    from copy import deepcopy

    ref_para = doc.paragraphs[idx]
    # Create a new paragraph element after the reference paragraph
    new_p = OxmlElement('w:p')
    # Copy paragraph properties from the reference if needed
    new_r = OxmlElement('w:r')
    new_t = OxmlElement('w:t')
    new_t.text = (
        "Figure 3: Watermark Robustness Dashboard - overview tab showing mean bit accuracy "
        "by transform for all three watermarking methods, and ensemble decision matrix "
        "colour-coded by best-performing method per cell."
    )
    new_r.append(new_t)
    new_p.append(new_r)
    ref_para._element.addnext(new_p)
    print(f"[Fix 4] Para {idx}: Dashboard screenshot placeholder added after.")
else:
    print("[Fix 4] ERROR: Could not find target text.")

# --- Fix 5: Update Section 2.3 to mention dashboard is live (after para 48) ---
idx = find_para_index(doc, "The full underlying datasets")
if idx is not None:
    ref_para = doc.paragraphs[idx]
    new_p = OxmlElement('w:p')
    new_r = OxmlElement('w:r')
    new_t = OxmlElement('w:t')
    new_t.text = (
        "The dashboard frontend is served by a FastAPI backend at http://127.0.0.1:8000, providing "
        "interactive visualisations of all benchmark results across four tabs: overview statistics, "
        "per-transform performance curves, ensemble decision matrix, and per-image drill-down."
    )
    new_r.append(new_t)
    new_p.append(new_r)
    ref_para._element.addnext(new_p)
    print(f"[Fix 5] Para {idx}: Dashboard live description added after.")
else:
    print("[Fix 5] ERROR: Could not find target text.")

# Save
doc.save(DOC_PATH)
print(f"\nDocument saved to {DOC_PATH}")
